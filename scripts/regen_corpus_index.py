# -*- coding: utf-8 -*-
"""重新生成 corpus_index.json - 修复 PDF/HTML 错乱问题"""
import os, json, re
from datetime import datetime

BASE = r'E:\desktop\火鹰合规\attrax\data\corpus'
OUT_JSON = os.path.join(BASE, 'corpus_index.json')
PROCESSED_DIR = os.path.join(BASE, 'processed')

# region 映射：子目录 -> region label
REGION_MAP = {
    'eu/regulations/pdfs': 'EU_Regulations_PDF',
    'eu/regulations/html': 'EU_Regulations_HTML',
    'eu/products': 'EU_Products',
    'eu/uk': 'UK',
    'us': 'US',
    'cn': 'CN',
    'asia/malaysia': 'MY',
    'asia/thailand': 'TH',
    'asia/singapore': 'SG',
    'asia/indonesia': 'ID',
    'asia/vietnam': 'VN',
    'middle_east/saudi': 'SA',
    'middle_east/uae': 'AE',
    'gcc': 'GCC',
    'intl/wipo': 'WIPO',
    'intl/un': 'UN',
    'reference': 'Reference',
    'screenshot_pending': 'Screenshot_Pending',
}

# ── 解析器 ────────────────────────────────────────────────────────────────────

def parse_html(fp):
    try:
        with open(fp, 'r', encoding='utf-8') as f:
            raw = f.read()
    except UnicodeDecodeError:
        with open(fp, 'r', encoding='gbk', errors='ignore') as f:
            raw = f.read()

    raw = re.sub(r'<script[^>]*>.*?</script>', '', raw, flags=re.DOTALL | re.IGNORECASE)
    raw = re.sub(r'<style[^>]*>.*?</style>', '', raw, flags=re.DOTALL | re.IGNORECASE)
    for noisy in ['nav', 'sidebar', 'footer', 'header', 'menu', 'advertisement']:
        raw = re.sub(rf'<[^>]+class="[^"]*{noisy}[^"]*"[^>]*>.*?</[^>]+>', '', raw, flags=re.DOTALL | re.IGNORECASE)

    text = re.sub(r'<[^>]+>', ' ', raw)
    text = re.sub(r'\s+', ' ', text).strip()

    tables = []
    for tm in re.findall(r'<table[^>]*>(.*?)</table>', raw, re.DOTALL | re.IGNORECASE):
        rows = []
        for tr in re.findall(r'<tr[^>]*>(.*?)</tr>', tm, re.DOTALL | re.IGNORECASE):
            cells = [re.sub(r'<[^>]+>', '', c).strip() for c in re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', tr, re.DOTALL | re.IGNORECASE)]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)

    return {
        'fileName': os.path.basename(fp),
        'sourcePath': fp,
        'sourceType': 'html',
        'pageCount': 1,
        'totalChars': len(text),
        'hasTables': len(tables) > 0,
        'tableCount': len(tables),
        'metadata': _extract_meta(text, 'html'),
    }

def parse_pdf(fp):
    import pdfplumber

    pages_data = []
    raw_text_parts = []
    all_tables = []
    total_chars = 0
    page_count = 0

    try:
        with pdfplumber.open(fp) as pdf:
            page_count = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ''
                raw_text_parts.append(text)
                total_chars += len(text)

                tables = []
                for t in page.extract_tables() or []:
                    cleaned = _clean_table(t)
                    if cleaned:
                        tables.append(cleaned)
                        all_tables.append(cleaned)

                pages_data.append({
                    'pageNumber': i + 1,
                    'charCount': len(text),
                    'hasText': len(text.strip()) > 50,
                    'hasTables': len(tables) > 0,
                    'textPreview': text[:200].replace('\n', ' ').strip() if text else '',
                })

    except Exception as e:
        return {'error': str(e), 'fileName': os.path.basename(fp), 'sourcePath': fp}

    raw_text = '\n\n'.join(raw_text_parts)

    return {
        'fileName': os.path.basename(fp),
        'sourcePath': fp,
        'sourceType': 'pdf',
        'pageCount': page_count,
        'totalChars': total_chars,
        'hasTables': len(all_tables) > 0,
        'tableCount': len(all_tables),
        'pages': pages_data,
        'rawText': raw_text,
        'metadata': _extract_meta(raw_text, 'pdf'),
    }

def parse_docx(fp):
    from docx import Document

    paragraphs_data = []
    all_tables = []
    total_chars = 0

    try:
        doc = Document(fp)
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                total_chars += len(text)
                paragraphs_data.append({'index': i, 'text': text})

        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(cells)
            if rows:
                all_tables.append(rows)

    except Exception as e:
        return {'error': str(e), 'fileName': os.path.basename(fp), 'sourcePath': fp}

    return {
        'fileName': os.path.basename(fp),
        'sourcePath': fp,
        'sourceType': 'docx',
        'pageCount': None,
        'paragraphCount': len(paragraphs_data),
        'totalChars': total_chars,
        'hasTables': len(all_tables) > 0,
        'tableCount': len(all_tables),
        'metadata': _extract_meta('\n'.join(p['text'] for p in paragraphs_data), 'docx'),
    }

def _clean_table(table):
    if not table or not table[0]:
        return None
    cleaned = []
    for row in table:
        if not row:
            continue
        cells = [str(c).strip() if c else '' for c in row]
        if any(c for c in cells if c):
            cleaned.append(cells)
    return cleaned if len(cleaned) > 1 else None

def _extract_meta(text, source_type):
    m = []

    if any(k in text for k in ['欧盟', 'EU', 'CE marking', 'GDPR', 'REACH', 'GPSR']):
        m.append('EU')
    if any(k in text for k in ['美国', 'FCC', 'CPSIA', 'COPPA', 'FDA', 'IRA']):
        m.append('US')
    if any(k in text for k in ['英国', 'UKCA', 'NSI Act']):
        m.append('UK')
    if any(k in text for k in ['中国', '出口管制', '商务部', 'GB', 'CCC']):
        m.append('CN')
    if any(k in text for k in ['新加坡', 'Singapore', 'PDPA', 'MAS']):
        m.append('SG')
    if any(k in text for k in ['越南', 'Vietnam']):
        m.append('VN')
    if any(k in text for k in ['马来西亚', 'Malaysia', 'PDPA 2010']):
        m.append('MY')
    if any(k in text for k in ['印尼', 'Indonesia', 'PDP Law']):
        m.append('ID')
    if any(k in text for k in ['沙特', 'Saudi', 'SABER', 'SFDA', 'PDPL']):
        m.append('SA')
    if any(k in text for k in ['阿联酋', 'UAE', 'Central Bank']):
        m.append('AE')
    if any(k in text for k in ['泰国', 'Thailand', 'PDPA 2019', 'B.E.']):
        m.append('TH')
    if any(k in text for k in ['G-Mark', 'Gulf', 'GCC', '海湾']):
        m.append('GCC')

    reg_types = []
    if any(k in text.lower() for k in ['product safety', '产品安全', 'GPSR', 'UL', 'IEC']):
        reg_types.append('product_safety')
    if any(k in text.lower() for k in ['chemical', 'REACH', 'RoHS', 'SVHC', '化学品']):
        reg_types.append('chemical')
    if any(k in text.lower() for k in ['data protection', 'GDPR', 'PDPA', '个人信息']):
        reg_types.append('data_privacy')
    if any(k in text.lower() for k in ['EMC', 'electromagnetic', '电磁兼容']):
        reg_types.append('EMC')
    if any(k in text.lower() for k in ['battery', 'lithium', '电池', '新电池法']):
        reg_types.append('battery')
    if any(k in text.lower() for k in ['export control', '出口管制', '两用物项']):
        reg_types.append('export_control')
    if any(k in text.lower() for k in ['investment', 'ODI', '境外投资', 'FIRRMA']):
        reg_types.append('investment')
    if any(k in text.lower() for k in ['toy', 'EN 71', '玩具']):
        reg_types.append('toy_safety')

    return {
        'detectedMarkets': m,
        'regulatoryTypes': reg_types,
        'charCount': len(text),
    }

def parse_file(fp):
    ext = os.path.splitext(fp)[1].lower()
    if ext in ('.pdf',):
        return parse_pdf(fp)
    elif ext in ('.html', '.htm'):
        return parse_html(fp)
    elif ext in ('.docx',):
        return parse_docx(fp)
    else:
        return {'error': f'unsupported: {ext}', 'fileName': os.path.basename(fp), 'sourcePath': fp}

# ── 主逻辑 ─────────────────────────────────────────────────────────────────────

# 确定每个文件的 region
def get_region(rel_path):
    for key in REGION_MAP:
        if rel_path.startswith(key + '/'):
            return key
    return None

# 扫描所有文件
all_files = []
for root, dirs, files in os.walk(BASE):
    # 跳过特定目录
    if any(x in root for x in ['processed', 'screenshot_pending']):
        continue
    for fn in files:
        if fn.startswith('.'):
            continue
        ext = os.path.splitext(fn)[1].lower()
        if ext not in ('.pdf', '.html', '.htm', '.docx'):
            continue
        fp = os.path.join(root, fn)
        rel = os.path.relpath(fp, BASE).replace('\\', '/')
        region_key = get_region(rel)
        if region_key is None:
            print(f'  SKIP (unknown region): {rel}')
            continue
        all_files.append({'fp': fp, 'rel': rel, 'region_key': region_key})

print(f'找到 {len(all_files)} 个文件待处理')

# 解析每个文件
all_entries = []
stats = {'total': 0, 'success': 0, 'failed': 0, 'by_region': {}}
by_region = {}

for f in sorted(all_files, key=lambda x: x['rel']):
    fp = fp0 = f['fp']
    region_key = f['region_key']
    rlabel = REGION_MAP[region_key]
    fn = os.path.basename(fp)

    stats['total'] += 1
    if region_key not in stats['by_region']:
        stats['by_region'][region_key] = {'total': 0, 'success': 0, 'failed': 0}
    stats['by_region'][region_key]['total'] += 1

    # safe output filename
    safe_name = re.sub(r'[^\w一-鿿\-]', '_', fn)
    safe_name = safe_name[:100]
    output_file = os.path.join(PROCESSED_DIR, f'{rlabel}_{safe_name}.json')

    print(f'  [{rlabel}] {fn}...', end=' ', flush=True)

    try:
        result = parse_file(fp)
        os.makedirs(PROCESSED_DIR, exist_ok=True)
        with open(output_file, 'w', encoding='utf-8') as fh:
            json.dump(result, fh, ensure_ascii=False, indent=2)

        entry = {
            'region': region_key,
            'regionLabel': rlabel,
            'fileName': fn,
            'sourcePath': fp,
            'outputPath': output_file,
            'sourceType': result.get('sourceType', 'unknown'),
            'pageCount': result.get('pageCount'),
            'paragraphCount': result.get('paragraphCount'),
            'totalChars': result.get('totalChars', 0),
            'hasTables': result.get('hasTables', False),
            'tableCount': result.get('tableCount', 0),
            'metadata': result.get('metadata', {}),
            'error': result.get('error'),
        }

        if 'error' not in result:
            stats['success'] += 1
            stats['by_region'][region_key]['success'] += 1
            print(f'OK ({result.get("totalChars", 0):,} chars)')
        else:
            stats['failed'] += 1
            stats['by_region'][region_key]['failed'] += 1
            print(f'FAILED: {result["error"]}')

        all_entries.append(entry)

    except Exception as e:
        stats['failed'] += 1
        stats['by_region'][region_key]['failed'] += 1
        print(f'ERROR: {e}')
        all_entries.append({
            'region': region_key,
            'regionLabel': rlabel,
            'fileName': fn,
            'sourcePath': fp,
            'error': str(e),
        })

# 生成索引
index = {
    'version': '1.0',
    'created': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
    'stats': stats,
    'files': all_entries,
}

with open(OUT_JSON, 'w', encoding='utf-8') as f:
    json.dump(index, f, ensure_ascii=False, indent=2)

print()
print('=' * 60)
print('  corpus_index.json 重新生成完成')
print('=' * 60)
print(f'  总文件: {stats["total"]}')
print(f'  成功:   {stats["success"]}')
print(f'  失败:   {stats["failed"]}')
print()
print('  按区域:')
for rk, rs in sorted(stats['by_region'].items()):
    ok_mark = 'OK' if rs['failed'] == 0 else 'PARTIAL'
    print(f'    [{rs["success"]}/{rs["total"]} {ok_mark}] {rk}')
print()
print(f'  输出目录: {PROCESSED_DIR}')
print(f'  索引文件: {OUT_JSON}')