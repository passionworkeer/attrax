from __future__ import annotations

import csv
from html import escape
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
except Exception:
    pass


def load_payload(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def build_compliance_lines(result: dict, locale: str) -> list[str]:
    if locale == "en":
        lines = [
            f"Product name: {result.get('productName') or 'Untitled product'}",
            f"Target markets: {' / '.join(result['targetMarkets'])}",
            f"Compliance score: {result['complianceScore']} / {result['scoreGrade']}",
            f"Generated at: {result['generatedAt']}",
            "",
            "Core findings:",
        ]
    else:
        lines = [
            f"产品名称: {result.get('productName') or '未命名产品'}",
            f"目标市场: {' / '.join(result['targetMarkets'])}",
            f"当前合规得分: {result['complianceScore']} / {result['scoreGrade']}",
            f"生成时间: {result['generatedAt']}",
            "",
            "核心结论:",
        ]
    for idx, risk in enumerate(result["riskPoints"], start=1):
        lines.append(f"{idx}. {risk['title']} - {risk['description']}")
    lines.extend(["", "Regulatory citations:" if locale == "en" else "重点法规引用:"])
    for risk in result["riskPoints"]:
        for regulation in risk["regulations"]:
            lines.append(
                f"- {regulation['market']} · {regulation['code']} · {regulation['name']}: {regulation['summary']}"
            )
    lines.extend(["", "Recommended remediation actions:" if locale == "en" else "推荐整改动作:"])
    for risk in result["riskPoints"]:
        if locale == "en":
            suffix = f" (Estimated {risk['estimatedFixCost']})" if risk.get("estimatedFixCost") else ""
        else:
            suffix = f" (预计 {risk['estimatedFixCost']})" if risk.get("estimatedFixCost") else ""
        lines.append(f"- {risk['title']}: {risk['recommendedAction']}{suffix}")
    return lines


def build_profit_lines(result: dict, locale: str) -> list[str]:
    critical = len([risk for risk in result["riskPoints"] if risk["severity"] == "critical"])
    warning = len([risk for risk in result["riskPoints"] if risk["severity"] == "warning"])
    if locale == "en":
        lines = [
            f"Product name: {result.get('productName') or 'Untitled product'}",
            f"Target markets: {' / '.join(result['targetMarkets'])}",
            f"Risk mix: {critical} critical / {warning} warning",
            "",
            "Decision summary:",
            f"- Current compliance score {result['complianceScore']} / {result['scoreGrade']}",
            "- Do not launch into target markets before the certification, label, and manual chain is closed",
            "- Finish the top 3 hotspot fixes before entering the formal certification stage",
            "",
            "Cost and profit hints:",
            "- Estimated heroic margin: ¥27 / unit",
            "- Real profit after compliance: ¥12 / unit",
            "- Per-platform compliance cost: ¥15 / unit",
            "- Estimated monthly loss: ¥12000",
            "",
            "Why remediation comes first:",
        ]
    else:
        lines = [
            f"产品名称: {result.get('productName') or '未命名产品'}",
            f"目标市场: {' / '.join(result['targetMarkets'])}",
            f"风险结构: 高危 {critical} 个 / 警告 {warning} 个",
            "",
            "决策摘要:",
            f"- 当前合规得分 {result['complianceScore']} / {result['scoreGrade']}",
            "- 在关键认证、标签和说明链路闭环前，不建议直接上架目标市场",
            "- 建议先完成前 3 个重点风险的整改后再进入正式认证阶段",
            "",
            "成本与利润提示:",
            "- 神勇出海预估利润: ¥27 / 件",
            "- 合规后真实利润: ¥12 / 件",
            "- 单平台合规成本: ¥15 / 件",
            "- 预估月损失利润: ¥12000",
            "",
            "为什么先整改:",
        ]
    for risk in result["riskPoints"]:
        if locale == "en":
            suffix = f", estimated {risk['estimatedFixCost']}" if risk.get("estimatedFixCost") else ""
        else:
            suffix = f"，预计成本 {risk['estimatedFixCost']}" if risk.get("estimatedFixCost") else ""
        lines.append(f"- {risk['title']}: {risk['recommendedAction']}{suffix}")
    return lines


def build_roadmap_rows(result: dict, locale: str) -> list[list[str]]:
    if locale == "en":
        return [
            ["Phase", "Action", "Output", "Note"],
            [
                "Document freeze",
                "Collect specification, BOM, nameplate, and supplier files",
                "Base product package",
                result.get("productName") or "",
            ],
            [
                "Label remediation",
                "Restore CE/UKCA, IO specs, and warning copy",
                "Shell and packaging artwork",
                " / ".join(result["targetMarkets"]),
            ],
            [
                "Risk review",
                "Confirm hotspot and citation closure",
                "Risk summary",
                f"{len(result['riskPoints'])} hotspots",
            ],
            [
                "Formal certification",
                "Enter lab testing and DoC flow",
                "Test and declaration files",
                "Recommended week 3-5",
            ],
            [
                "Listing review",
                "Align listing, hero image, and manual",
                "Launch package",
                "Only enter the market after closure",
            ],
        ]

    return [
        ["阶段", "动作", "输出", "备注"],
        ["资料冻结", "整理规格书、BOM、铭牌与供应商资料", "产品基础资料包", result.get("productName") or ""],
        ["标签整改", "补齐 CE/UKCA、输入输出规格和警示语", "外壳与包装图稿", " / ".join(result["targetMarkets"])],
        ["风险复核", "确认风险点与法规引用是否闭环", "风险总表", f"{len(result['riskPoints'])} 个热点"],
        ["正式认证", "进入实验室测试与 DoC 流程", "测试与声明文件", "建议第 3-5 周"],
        ["上架复核", "同步 Listing / 主图 / 说明书", "上架资料", "完成后再进目标市场"],
    ]


def write_csv(output_path: Path, rows: list[list[str]]) -> None:
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)


def pdf_font(locale: str) -> str:
    return "Helvetica" if locale == "en" else "STSong-Light"


def pdf_bold_font(locale: str) -> str:
    return "Helvetica-Bold" if locale == "en" else "STSong-Light"


def as_paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)


def metric_cards(result: dict, locale: str) -> list[list[str]]:
    financial = result.get("financialSummary", {})
    if locale == "en":
        return [
            ["Compliance score", f"{result['complianceScore']} / {result['scoreGrade']}"],
            ["True net profit", financial.get("trueNetProfit", "-")],
            ["Compliance cost", financial.get("complianceCost", "-")],
            ["Monthly exposure", financial.get("monthlyNetProfit", "-")],
        ]

    return [
        ["合规得分", f"{result['complianceScore']} / {result['scoreGrade']}"],
        ["真实净利", financial.get("trueNetProfit", "-")],
        ["合规成本", financial.get("complianceCost", "-")],
        ["月度风险敞口", financial.get("monthlyNetProfit", "-")],
    ]


def set_docx_table_font(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(9)


def add_docx_metric_table(document: Document, result: dict, locale: str) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in metric_cards(result, locale):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    set_docx_table_font(table)


def add_docx_risk_table(document: Document, result: dict, locale: str) -> None:
    headers = (
        ["Risk", "Severity", "Confidence", "Fix cost", "Recommended action"]
        if locale == "en"
        else ["风险点", "级别", "置信度", "整改成本", "建议动作"]
    )
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    for risk in result["riskPoints"]:
        cells = table.add_row().cells
        cells[0].text = risk["title"]
        cells[1].text = risk["severity"]
        cells[2].text = f"{round(risk['confidence'] * 100)}%"
        cells[3].text = risk.get("estimatedFixCost") or "-"
        cells[4].text = risk["recommendedAction"]
    set_docx_table_font(table)


def build_branded_pdf(output_path: Path, title: str, result: dict, report_type: str, locale: str) -> None:
    font = pdf_font(locale)
    bold_font = pdf_bold_font(locale)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "blaze-title",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=22,
        leading=29,
        textColor=colors.white,
        spaceAfter=6,
    )
    heading_style = ParagraphStyle(
        "blaze-heading",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=12,
        leading=18,
        textColor=colors.HexColor("#1f315c"),
        spaceAfter=8,
        spaceBefore=10,
    )
    body_style = ParagraphStyle(
        "blaze-body",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=10,
        leading=15,
        textColor=colors.HexColor("#33435f"),
        spaceAfter=4,
    )
    small_style = ParagraphStyle(
        "blaze-small",
        parent=body_style,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#65738f"),
    )

    product_name = result.get("productName") or ("Untitled product" if locale == "en" else "未命名产品")
    market_label = " / ".join(result["targetMarkets"])
    generated_label = "Generated at" if locale == "en" else "生成时间"
    market_title = "Target markets" if locale == "en" else "目标市场"

    header = Table(
        [
            [
                as_paragraph("BLAZE HAWKS", ParagraphStyle("brand", parent=body_style, fontName=bold_font, fontSize=10, leading=13, textColor=colors.HexColor("#ff9e3d"))),
                as_paragraph(generated_label, ParagraphStyle("date-label", parent=small_style, textColor=colors.HexColor("#b6c5df"))),
            ],
            [
                as_paragraph(title, title_style),
                as_paragraph(result["generatedAt"][:10], ParagraphStyle("date-value", parent=body_style, fontName=bold_font, alignment=2, textColor=colors.white)),
            ],
            [
                as_paragraph(f"{product_name} · {market_title}: {market_label}", ParagraphStyle("sub", parent=body_style, fontName=font, textColor=colors.HexColor("#dfe9ff"))),
                "",
            ],
        ],
        colWidths=[360, 120],
    )
    header.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#101b31")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#22365f")),
                ("SPAN", (0, 2), (-1, 2)),
                ("LEFTPADDING", (0, 0), (-1, -1), 18),
                ("RIGHTPADDING", (0, 0), (-1, -1), 18),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    metric_style = ParagraphStyle(
        "metric-card",
        parent=body_style,
        fontName=bold_font,
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#1f315c"),
        alignment=1,
    )
    metric_table = Table(
        [[as_paragraph(f"{label}\n{value}", metric_style) for label, value in metric_cards(result, locale)]],
        colWidths=[120, 120, 120, 120],
    )
    metric_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f3f7fc")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#d9e5f6")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9e5f6")),
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#1f315c")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )

    risk_headers = (
        ["Risk", "Severity", "Confidence", "Fix cost", "Recommended action"]
        if locale == "en"
        else ["风险点", "级别", "置信度", "整改成本", "建议动作"]
    )
    risk_rows = [[as_paragraph(cell, ParagraphStyle("risk-head", parent=body_style, fontName=bold_font, textColor=colors.HexColor("#182848"))) for cell in risk_headers]]
    for risk in result["riskPoints"]:
        risk_rows.append(
            [
                as_paragraph(risk["title"], body_style),
                as_paragraph(risk["severity"], body_style),
                as_paragraph(f"{round(risk['confidence'] * 100)}%", body_style),
                as_paragraph(risk.get("estimatedFixCost") or "-", body_style),
                as_paragraph(risk["recommendedAction"], body_style),
            ]
        )

    risk_table = Table(risk_rows, repeatRows=1, colWidths=[92, 58, 62, 72, 196])
    risk_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf2fc")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
                ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#d7e4f5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )

    if report_type == "profit":
        summary_title = "AI decision" if locale == "en" else "AI 决策摘要"
        summary_lines = (
            [
                "Do not launch before the certification, label, and manual chain is closed.",
                "Use the compliance cost as the floor for the first-batch pricing decision.",
            ]
            if locale == "en"
            else [
                "关键认证、标签与说明链路闭环前，不建议直接上架目标市场。",
                "首批试销前，应把合规成本作为报价底线并预留整改周期。",
            ]
        )
    else:
        summary_title = "Compliance summary" if locale == "en" else "合规摘要"
        summary_lines = (
            [
                f"{len(result['riskPoints'])} visual hotspots are mapped to market-specific citations.",
                "The first remediation pass should close marks, specifications, and warning copy.",
            ]
            if locale == "en"
            else [
                f"已将 {len(result['riskPoints'])} 个视觉风险点映射到目标市场法规引用。",
                "第一轮整改应优先闭环认证标识、输入输出规格与多语言警示。",
            ]
        )

    story = [
        header,
        Spacer(1, 16),
        metric_table,
        Spacer(1, 14),
        as_paragraph(summary_title, heading_style),
        *[as_paragraph(f"• {line}", body_style) for line in summary_lines],
        Spacer(1, 10),
        as_paragraph("Risk hotspot table" if locale == "en" else "风险热点清单", heading_style),
        risk_table,
    ]

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=42,
        rightMargin=42,
        topMargin=42,
        bottomMargin=42,
    )
    document.build(story)


def build_branded_docx(output_path: Path, title: str, result: dict, report_type: str, locale: str) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10)

    title_paragraph = document.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title_paragraph.runs:
        run.font.color.rgb = RGBColor(19, 36, 73)

    product_name = result.get("productName") or ("Untitled product" if locale == "en" else "未命名产品")
    document.add_paragraph(
        f"{'Product name' if locale == 'en' else '产品名称'}: {product_name}"
    )
    document.add_paragraph(
        f"{'Target markets' if locale == 'en' else '目标市场'}: {' / '.join(result['targetMarkets'])}"
    )
    document.add_paragraph(
        f"{'Generated at' if locale == 'en' else '生成时间'}: {result['generatedAt']}"
    )

    document.add_heading("Metrics" if locale == "en" else "关键指标", level=1)
    add_docx_metric_table(document, result, locale)

    document.add_heading("Risk hotspot table" if locale == "en" else "风险热点清单", level=1)
    add_docx_risk_table(document, result, locale)

    document.add_heading("Recommended next steps" if locale == "en" else "推荐下一步", level=1)
    if report_type == "profit":
        steps = (
            [
                "Hold launch until the certification, label, and manual chain is closed.",
                "Use compliance cost and expected exposure as the pricing floor.",
                "Re-run the scan after the first remediation pass.",
            ]
            if locale == "en"
            else [
                "认证、标签与说明链路闭环前，暂缓直接进入目标市场。",
                "以合规成本和月度风险敞口作为首批报价底线。",
                "完成第一轮整改后重新扫描并更新报告。",
            ]
        )
    else:
        steps = (
            [
                "Close the top certification, spec, and warning hotspots first.",
                "Align listing imagery, manual, and supplier evidence before lab work.",
                "Export the roadmap for supplier and legal review.",
            ]
            if locale == "en"
            else [
                "优先闭环认证标识、规格参数与警示语风险点。",
                "进入实验室前，同步 Listing 主图、说明书与供应商证据。",
                "导出路线图，供供应商与法务共同复核。",
            ]
        )
    for step in steps:
        document.add_paragraph(step, style="List Bullet")
    document.save(output_path)


def build_roadmap_pdf(output_path: Path, result: dict, locale: str) -> None:
    rows = build_roadmap_rows(result, locale)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "roadmap-title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=26,
        textColor=colors.HexColor("#132449"),
        spaceAfter=16,
    )
    body_style = ParagraphStyle(
        "roadmap-body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=colors.HexColor("#33435f"),
        spaceAfter=4,
    )

    table = Table(rows, repeatRows=1, colWidths=[72, 144, 126, 110])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf2fc")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1d3158")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d5e4f9")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEADING", (0, 0), (-1, -1), 12),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fbff")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )

    story = [
        Paragraph(
            "CompliPilot · Compliance Roadmap Report"
            if locale == "en"
            else "规航AI · 合规路线图报告",
            title_style,
        ),
        Paragraph(
            f"{'Product name' if locale == 'en' else '产品名称'}: {result.get('productName') or ('Untitled product' if locale == 'en' else '未命名产品')}",
            body_style,
        ),
        Paragraph(
            f"{'Target markets' if locale == 'en' else '目标市场'}: {' / '.join(result['targetMarkets'])}",
            body_style,
        ),
        Spacer(1, 12),
        table,
    ]
    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    document.build(story)


def build_roadmap_docx(output_path: Path, result: dict, locale: str) -> None:
    document = Document()
    document.add_heading(
        "CompliPilot · Compliance Roadmap Report"
        if locale == "en"
        else "规航AI · 合规路线图报告",
        level=0,
    )
    document.add_paragraph(
        f"{'Product name' if locale == 'en' else '产品名称'}: {result.get('productName') or ('Untitled product' if locale == 'en' else '未命名产品')}"
    )
    document.add_paragraph(
        f"{'Target markets' if locale == 'en' else '目标市场'}: {' / '.join(result['targetMarkets'])}"
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = build_roadmap_rows(result, locale)[0]
    for index, label in enumerate(headers):
        table.rows[0].cells[index].text = label
    for row in build_roadmap_rows(result, locale)[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.save(output_path)


def main() -> int:
    if len(sys.argv) != 5:
        raise SystemExit("usage: generate_report_file.py <input-json> <output-path> <report-type> <format>")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    report_type = sys.argv[3]
    output_format = sys.argv[4]

    payload = load_payload(input_path)
    locale = payload.get("locale", "zh")
    result = payload["result"]

    if report_type == "compliance":
        if output_format == "pdf":
            build_branded_pdf(
                output_path,
                "CompliPilot · Compliance Scan Report" if locale == "en" else "规航AI · 合规扫描报告",
                result,
                report_type,
                locale,
            )
        elif output_format == "docx":
            build_branded_docx(
                output_path,
                "CompliPilot · Compliance Scan Report" if locale == "en" else "规航AI · 合规扫描报告",
                result,
                report_type,
                locale,
            )
        else:
            raise SystemExit("unsupported format")
    elif report_type == "profit":
        if output_format == "pdf":
            build_branded_pdf(
                output_path,
                "CompliPilot · Compliance Cost Impact / AI Decision Report"
                if locale == "en"
                else "规航AI · 合规成本影响 / AI 决策报告",
                result,
                report_type,
                locale,
            )
        elif output_format == "docx":
            build_branded_docx(
                output_path,
                "CompliPilot · Compliance Cost Impact / AI Decision Report"
                if locale == "en"
                else "规航AI · 合规成本影响 / AI 决策报告",
                result,
                report_type,
                locale,
            )
        else:
            raise SystemExit("unsupported format")
    elif report_type == "roadmap":
        if output_format == "pdf":
            build_roadmap_pdf(output_path, result, locale)
        elif output_format == "docx":
            build_roadmap_docx(output_path, result, locale)
        else:
            raise SystemExit("unsupported format")
    else:
        raise SystemExit("unsupported report type")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())

